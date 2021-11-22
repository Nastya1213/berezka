from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference

wb = load_workbook('Result.xlsx')
# получаем лист, с которым будем работать
sheet = wb['Первый лист']
f_name = sheet['F1'].value
# создаем диаграмму
chart = BarChart()
chart.title = 'Результы тестов'
data = Reference(sheet, min_col=2, min_row=2, max_col=2, max_row=4)
chart.add_data(data)
# добавляем диаграмму на лист
sheet.add_chart(chart, 'G2')
wb.save(f'{f_name}.xlsx')





